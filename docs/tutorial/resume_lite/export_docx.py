"""Step 3 — a Word document with python-docx."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from model import Profile


def period(start: str, end: str) -> str:
    return f"{start} – {end or 'Present'}"


def export_docx(profile: Profile, path: Path) -> Path:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(profile.name, level=0)
    if profile.headline:
        doc.add_paragraph(profile.headline)
    if profile.email:
        doc.add_paragraph(profile.email)
    if profile.summary:
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(profile.summary)
    if profile.jobs:
        doc.add_heading("Experience", level=1)
        for job in profile.jobs:
            p = doc.add_paragraph()
            p.add_run(f"{job.title}, {job.organisation}").bold = True
            p.add_run(f"    {period(job.start, job.end)}")
            for a in job.achievements:
                doc.add_paragraph(a, style="List Bullet")
    if profile.skills:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(profile.skills))
    doc.save(path)
    return path
