"""DOCX renderer: LayoutDocument → genuine, editable Word document.

Uses real Word structures: built-in Title / Heading 1 / Heading 2 / Normal
styles (configured to the template), a numbering definition for bullets,
tab stops for right-aligned dates and hanging gutters, paragraph borders and
shading, tables, external hyperlinks, PAGE/NUMPAGES fields in the footer, and
core document properties.
"""

from __future__ import annotations

import re
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ris.rendering.blocks import (
    BannerBlock,
    Border,
    ColumnsBlock,
    ImageBlock,
    LabeledSection,
    LayoutDocument,
    PageBreakBlock,
    Para,
    ParaStyle,
    Rule,
    Run,
    Spacer,
    TableBlock,
)

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")
BULLET_NUM_BASE = 90


def clean_text(text: str) -> str:
    return _ILLEGAL_XML.sub("", text or "")


def _twips(pt: float) -> int:
    return int(round(pt * 20))


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


def _set_rfonts(rpr, font: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)


def _set_lang(rpr, lang: str) -> None:
    el = rpr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rpr.append(el)
    el.set(qn("w:val"), lang)


def _set_spacing(rpr, pts: float) -> None:
    el = rpr.find(qn("w:spacing"))
    if el is None:
        el = OxmlElement("w:spacing")
        rpr.append(el)
    el.set(qn("w:val"), str(_twips(pts)))


def _border_el(tag: str, b: Border, val: str = "single"):
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(max(2, int(round(b.width * 8)))))
    el.set(qn("w:space"), str(int(round(b.space))))
    el.set(qn("w:color"), b.color.lstrip("#").upper())
    return el


def _ppr_borders(paragraph, top: Border | None = None, bottom: Border | None = None, left: Border | None = None) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:pBdr"))
    if old is not None:
        ppr.remove(old)
    pbdr = OxmlElement("w:pBdr")
    if top:
        pbdr.append(_border_el("top", top))
    if left:
        pbdr.append(_border_el("left", left))
    if bottom:
        pbdr.append(_border_el("bottom", bottom))
    _insert_ppr_child(ppr, pbdr)


# Order of pPr children required by the WordprocessingML schema (subset we use).
_PPR_ORDER = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:widowControl", "w:numPr", "w:pBdr",
              "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:outlineLvl", "w:rPr"]


def _insert_ppr_child(ppr, child) -> None:
    tag = child.tag
    names = [qn(n) for n in _PPR_ORDER]
    idx = names.index(tag) if tag in names else len(names)
    for existing in list(ppr):
        if existing.tag in names and names.index(existing.tag) > idx:
            existing.addprevious(child)
            return
    ppr.append(child)


def _ppr_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#").upper())
    _insert_ppr_child(ppr, shd)


def _cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#").upper())
    tcpr.append(shd)


class DocxRenderer:
    def __init__(self, layout: LayoutDocument) -> None:
        self.layout = layout
        self.doc = Document()
        self.style_names: dict[str, str] = {}
        self._bullet_nums: dict[tuple[str, int], int] = {}
        self._next_num = BULLET_NUM_BASE

    # ------------------------------------------------------------------ setup
    def _setup_page(self) -> None:
        pg = self.layout.page
        sec = self.doc.sections[0]
        sec.page_width = Pt(pg.width)
        sec.page_height = Pt(pg.height)
        sec.top_margin = Pt(pg.margin_top)
        sec.bottom_margin = Pt(pg.margin_bottom)
        sec.left_margin = Pt(pg.margin_left)
        sec.right_margin = Pt(pg.margin_right)
        sec.header_distance = Pt(min(pg.margin_top * 0.5, 28))
        sec.footer_distance = Pt(min(pg.margin_bottom * 0.45, 24))
        sec.start_type = WD_SECTION.NEW_PAGE

    def _configure_style(self, style, ps: ParaStyle) -> None:
        font = style.font
        font.name = ps.font
        font.size = Pt(ps.size)
        font.bold = ps.bold
        font.italic = ps.italic
        font.small_caps = ps.small_caps
        font.all_caps = False
        font.underline = False
        font.color.rgb = _rgb(ps.color)
        rpr = style.element.get_or_add_rPr()
        _set_rfonts(rpr, ps.font)
        _set_lang(rpr, self.layout.lang)
        if ps.letter_spacing:
            _set_spacing(rpr, ps.letter_spacing)
        pf = style.paragraph_format
        pf.space_before = Pt(ps.space_before)
        pf.space_after = Pt(ps.space_after)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = ps.line_spacing
        pf.alignment = _ALIGN.get(ps.align, WD_ALIGN_PARAGRAPH.LEFT)
        pf.left_indent = Pt(ps.left_indent) if ps.left_indent else Pt(0)
        pf.first_line_indent = Pt(ps.first_line_indent) if ps.first_line_indent else Pt(0)
        pf.keep_with_next = ps.keep_with_next
        pf.widow_control = True
        ppr = style.element.get_or_add_pPr()
        for tag in ("w:pBdr", "w:shd"):
            old = ppr.find(qn(tag))
            if old is not None:
                ppr.remove(old)
        if ps.outline_level is not None:
            ol = ppr.find(qn("w:outlineLvl"))
            if ol is None:
                ol = OxmlElement("w:outlineLvl")
                _insert_ppr_child(ppr, ol)
            ol.set(qn("w:val"), str(ps.outline_level))

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        # Document defaults so unstyled text matches the template.
        normal = self.layout.styles["Normal"]
        defaults = self.doc.styles.element.find(qn("w:docDefaults"))
        if defaults is not None:
            rpr_default = defaults.find(qn("w:rPrDefault"))
            if rpr_default is not None:
                rpr = rpr_default.find(qn("w:rPr"))
                if rpr is not None:
                    _set_rfonts(rpr, normal.font)
                    _set_lang(rpr, self.layout.lang)
        claimed: set[str] = set()
        from docx.enum.style import WD_STYLE_TYPE

        for name, ps in self.layout.styles.items():
            if name == "Normal":
                self._configure_style(styles["Normal"], ps)
                self.style_names[name] = "Normal"
                continue
            if ps.docx_base and ps.docx_base not in claimed and ps.docx_base in [s.name for s in styles]:
                st = styles[ps.docx_base]
                claimed.add(ps.docx_base)
            else:
                spaced = re.sub(r"(?<!^)([A-Z])", r" \1", name)
                st_name = "RIS " + spaced
                st = styles.add_style(st_name, WD_STYLE_TYPE.PARAGRAPH)
                st.base_style = styles["Normal"]
                st.quick_style = name in ("Body", "Bullet", "EntrySub")
            self._configure_style(st, ps)
            self.style_names[name] = st.name
        # Hyperlink character style colour.
        try:
            hl = styles["Hyperlink"]
            hl.font.color.rgb = _rgb(self.layout.link_color)
            hl.font.underline = False
        except KeyError:
            pass

    # ------------------------------------------------------------------ numbering (bullets)
    def _numbering_root(self):
        return self.doc.part.numbering_part.element

    def _bullet_num_id(self, glyph: str, font: str) -> int:
        key = (glyph, 0)
        if key in self._bullet_nums:
            return self._bullet_nums[key]
        root = self._numbering_root()
        abs_id = self._next_num
        num_id = self._next_num
        self._next_num += 1
        ind = self.layout.bullet_indent
        absn = OxmlElement("w:abstractNum")
        absn.set(qn("w:abstractNumId"), str(abs_id))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "hybridMultilevel")
        absn.append(mlt)
        glyphs = [glyph, "–", "·"]
        for lvl in range(3):
            el = OxmlElement("w:lvl")
            el.set(qn("w:ilvl"), str(lvl))
            for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", glyphs[lvl]), ("w:lvlJc", "left")):
                child = OxmlElement(tag)
                child.set(qn("w:val"), val)
                el.append(child)
            ppr = OxmlElement("w:pPr")
            indel = OxmlElement("w:ind")
            indel.set(qn("w:left"), str(_twips(ind * (lvl + 1))))
            indel.set(qn("w:hanging"), str(_twips(ind)))
            ppr.append(indel)
            el.append(ppr)
            rpr = OxmlElement("w:rPr")
            _set_rfonts(rpr, font)
            el.append(rpr)
            absn.append(el)
        first_num = root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(absn)
        else:
            root.append(absn)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        an = OxmlElement("w:abstractNumId")
        an.set(qn("w:val"), str(abs_id))
        num.append(an)
        root.append(num)
        self._bullet_nums[key] = num_id
        return num_id

    # ------------------------------------------------------------------ runs and paragraphs
    def _add_run(self, paragraph, r: Run) -> None:
        text = clean_text(r.text)
        run = paragraph.add_run(text)
        if r.bold:
            run.bold = True
        if r.italic:
            run.italic = True
        if r.size:
            run.font.size = Pt(r.size)
        if r.small_caps:
            run.font.small_caps = True
        if r.underline:
            run.underline = True
        if r.font:
            _set_rfonts(run._r.get_or_add_rPr(), r.font)
        color = r.color or (self.layout.link_color if r.link else None)
        if color:
            run.font.color.rgb = _rgb(color)
        if r.link:
            r_id = paragraph.part.relate_to(r.link, RT.HYPERLINK, is_external=True)
            hl = OxmlElement("w:hyperlink")
            hl.set(qn("r:id"), r_id)
            hl.set(qn("w:history"), "1")
            run._r.addprevious(hl)
            hl.append(run._r)

    def _render_para(self, container, p: Para, width: float, reuse=None):
        style_name = self.style_names.get(p.style, "Normal")
        ps = self.layout.styles.get(p.style, self.layout.styles["Normal"])
        if reuse is not None:
            para = reuse
            para.style = self.doc.styles[style_name]
        else:
            para = container.add_paragraph(style=style_name)
        pf = para.paragraph_format
        if p.align:
            pf.alignment = _ALIGN.get(p.align, WD_ALIGN_PARAGRAPH.LEFT)
        if p.keep_with_next is not None:
            pf.keep_with_next = p.keep_with_next
        if p.space_before is not None:
            pf.space_before = Pt(p.space_before)
        if p.space_after is not None:
            pf.space_after = Pt(p.space_after)
        base_indent = p.left_indent if p.left_indent is not None else ps.left_indent
        if p.gutter_runs is not None:
            pf.left_indent = Pt(base_indent)
            pf.first_line_indent = Pt(-p.gutter_width)
            pf.tab_stops.add_tab_stop(Pt(base_indent))
        elif p.bullet:
            num_id = self._bullet_num_id(p.bullet, ps.font)
            ppr = para._p.get_or_add_pPr()
            numpr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), str(min(p.bullet_level, 2)))
            numid = OxmlElement("w:numId")
            numid.set(qn("w:val"), str(num_id))
            numpr.append(ilvl)
            numpr.append(numid)
            _insert_ppr_child(ppr, numpr)
            indent = base_indent if p.left_indent is not None else self.layout.bullet_indent * (p.bullet_level + 1)
            pf.left_indent = Pt(indent)
            pf.first_line_indent = Pt(-self.layout.bullet_indent)
        elif p.left_indent is not None:
            pf.left_indent = Pt(p.left_indent)
        if p.right_runs:
            pf.tab_stops.add_tab_stop(Pt(max(width - 0.5, 10)), WD_TAB_ALIGNMENT.RIGHT)
        if p.border_bottom or p.border_left or p.border_top:
            _ppr_borders(para, top=p.border_top, bottom=p.border_bottom, left=p.border_left)
        if p.shading:
            _ppr_shading(para, p.shading)
        if p.gutter_runs is not None:
            for r in p.gutter_runs:
                self._add_run(para, r)
            para.add_run("\t")
        for r in p.runs:
            self._add_run(para, r)
        if p.right_runs:
            para.add_run("\t")
            for r in p.right_runs:
                self._add_run(para, r)
        return para

    # ------------------------------------------------------------------ containers
    def _table(self, container, rows: int, cols: int, widths_pt: list[float], borders: str,
               border_color: str = "#000000", border_width: float = 0.5, padding: float = 3.0):
        table = container.add_table(rows=rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        tbl = table._tbl
        tblpr = tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblpr.append(layout)
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW")
            tblpr.append(tblw)
        tblw.set(qn("w:w"), str(_twips(sum(widths_pt))))
        tblw.set(qn("w:type"), "dxa")
        bdr = OxmlElement("w:tblBorders")
        color = border_color.lstrip("#").upper()
        sz = str(max(2, int(round(border_width * 8))))
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            on = borders == "grid" or (borders == "rows" and edge in ("bottom", "insideH")) or (
                borders == "header_rule" and edge == "insideH")
            el.set(qn("w:val"), "single" if on else "nil")
            if on:
                el.set(qn("w:sz"), sz)
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), color)
            bdr.append(el)
        tblpr.append(bdr)
        mar = OxmlElement("w:tblCellMar")
        for edge, val in (("top", padding * 0.6), ("left", padding), ("bottom", padding * 0.6), ("right", padding)):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:w"), str(_twips(val)))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tblpr.append(mar)
        grid = tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(_twips(widths_pt[i])))
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Pt(widths_pt[i])
        return table

    def _fill_cell(self, cell, blocks: list, width: float) -> None:
        first = cell.paragraphs[0] if cell.paragraphs else None
        used_first = False
        for b in blocks:
            if isinstance(b, Para) and first is not None and not used_first:
                self._render_para(cell, b, width, reuse=first)
                used_first = True
            else:
                self._render_block(cell, b, width)
                used_first = True
        if not blocks and first is not None:
            first.paragraph_format.space_after = Pt(0)

    def _render_block(self, container, b, width: float) -> None:
        if isinstance(b, Para):
            self._render_para(container, b, width)
        elif isinstance(b, Spacer):
            if b.height <= 0:
                return
            para = container.add_paragraph(style=self.style_names["Normal"])
            pf = para.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(max(b.height, 1))
            run = para.add_run("")
            run.font.size = Pt(1)
        elif isinstance(b, Rule):
            para = container.add_paragraph(style=self.style_names["Normal"])
            pf = para.paragraph_format
            pf.space_before = Pt(b.space_before)
            pf.space_after = Pt(b.space_after)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(1)
            _ppr_borders(para, bottom=Border(b.color, b.width, 0))
        elif isinstance(b, PageBreakBlock):
            para = container.add_paragraph()
            para.add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(b, ImageBlock):
            para = container.add_paragraph(style=self.style_names["Normal"])
            para.paragraph_format.alignment = _ALIGN.get(b.align, WD_ALIGN_PARAGRAPH.RIGHT)
            para.paragraph_format.space_after = Pt(b.space_after)
            para.paragraph_format.keep_with_next = b.keep_with_next
            try:
                shape = para.add_run().add_picture(b.source(), width=Pt(min(b.width, width)))
                if b.alt:
                    shape._inline.docPr.set("descr", b.alt[:1000])
                    shape._inline.docPr.set("title", b.alt.split(":")[0][:120])
            except Exception:  # unreadable image: leave the paragraph empty, validation reports it
                para.add_run("")
        elif isinstance(b, TableBlock):
            widths = [width * f for f in b.col_widths]
            table = self._table(container, len(b.rows), len(b.col_widths), widths, b.borders, b.border_color,
                                b.border_width, b.padding)
            for ri, row in enumerate(b.rows):
                if ri < b.header_rows:
                    trpr = table.rows[ri]._tr.get_or_add_trPr()
                    hdr = OxmlElement("w:tblHeader")
                    trpr.append(hdr)
                for ci, cell_spec in enumerate(row):
                    cell = table.cell(ri, ci)
                    fill = cell_spec.fill or (b.header_fill if ri < b.header_rows else None)
                    if fill:
                        _cell_shading(cell, fill)
                    self._fill_cell(cell, cell_spec.blocks, widths[ci] - 2 * b.padding)
            self._after_table(container)
        elif isinstance(b, ColumnsBlock):
            lw = width * b.left_frac
            rw = width - lw
            table = self._table(container, 1, 2, [lw, rw], "none", padding=0)
            left, right = table.cell(0, 0), table.cell(0, 1)
            self._cell_margins(left, b.padding if b.left_fill else 0, b.gap / 2 + (b.padding if b.left_fill else 0))
            self._cell_margins(right, b.gap / 2 + (b.padding if b.right_fill else 0), b.padding if b.right_fill else 0)
            if b.left_fill:
                _cell_shading(left, b.left_fill)
            if b.right_fill:
                _cell_shading(right, b.right_fill)
            lpad = (b.padding if b.left_fill else 0) + b.gap / 2 + (b.padding if b.left_fill else 0)
            rpad = b.gap / 2 + 2 * (b.padding if b.right_fill else 0)
            self._fill_cell(left, b.left, lw - lpad)
            self._fill_cell(right, b.right, rw - rpad)
            self._after_table(container)
        elif isinstance(b, BannerBlock):
            table = self._table(container, 1, 1, [width], "none", padding=b.padding)
            cell = table.cell(0, 0)
            _cell_shading(cell, b.fill)
            self._fill_cell(cell, b.blocks, width - 2 * b.padding)
            self._after_table(container)
        elif isinstance(b, LabeledSection):
            lw = b.label_width
            rw = width - lw
            table = self._table(container, 1, 2, [lw, rw], "none", padding=0)
            self._cell_margins(table.cell(0, 0), 0, b.gap)
            self._fill_cell(table.cell(0, 0), [b.label], lw - b.gap)
            self._fill_cell(table.cell(0, 1), b.blocks, rw)
            self._after_table(container)

    def _cell_margins(self, cell, left: float, right: float) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for edge, val in (("left", left), ("right", right), ("top", 0), ("bottom", 0)):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:w"), str(_twips(val)))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tcpr.append(mar)

    def _after_table(self, container) -> None:
        # Keep a tiny paragraph after tables so consecutive tables do not merge in Word.
        para = container.add_paragraph(style=self.style_names["Normal"])
        pf = para.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(2)

    # ------------------------------------------------------------------ footer
    def _add_field(self, paragraph, instr: str, placeholder: str = "1") -> None:
        def fld(kind: str):
            r = paragraph.add_run()
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
            return r

        fld("begin")
        r = paragraph.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        r._r.append(it)
        fld("separate")
        paragraph.add_run(placeholder)
        fld("end")

    def _setup_footer(self) -> None:
        ft = self.layout.footer
        if ft.page_numbers == "none" and not ft.left_text and not ft.note:
            return
        sec = self.doc.sections[0]
        footer = sec.footer
        para = footer.paragraphs[0]
        width = self.layout.page.text_width
        pf = para.paragraph_format
        pf.tab_stops.add_tab_stop(Pt(width), WD_TAB_ALIGNMENT.RIGHT)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)

        def styled(run):
            run.font.size = Pt(ft.size)
            run.font.color.rgb = _rgb(ft.color)
            run.font.name = ft.font
            _set_rfonts(run._r.get_or_add_rPr(), ft.font)
            return run

        if ft.left_text:
            styled(para.add_run(clean_text(ft.left_text)))
        if ft.page_numbers != "none":
            styled(para.add_run("\t"))
            if ft.page_numbers == "page_x_of_y":
                styled(para.add_run("Page "))
            self._add_field(para, "PAGE")
            if ft.page_numbers in ("page_x_of_y", "x_of_y"):
                styled(para.add_run(" of " if ft.page_numbers == "page_x_of_y" else " / "))
                self._add_field(para, "NUMPAGES")
            for r in para.runs:
                styled(r)
        if ft.note:
            note = footer.add_paragraph()
            note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            note.paragraph_format.space_before = Pt(1)
            styled(note.add_run(clean_text(ft.note))).italic = True

    def _core_properties(self) -> None:
        md = self.layout.metadata
        cp = self.doc.core_properties
        now = datetime.now(UTC).replace(microsecond=0, tzinfo=None)
        cp.author = md.get("author", "")
        cp.last_modified_by = md.get("author", "")
        cp.title = md.get("title", "")
        cp.subject = md.get("subject", "")
        cp.keywords = md.get("keywords", "")
        cp.comments = md.get("creator", "Resume Intelligence Studio")
        cp.category = "Resume"
        cp.created = now
        cp.modified = now
        cp.revision = 1
        cp.language = self.layout.lang

    # ------------------------------------------------------------------ main
    def render(self, path: Path) -> Path:
        self._setup_page()
        self._setup_styles()
        width = self.layout.page.text_width
        for b in self.layout.blocks:
            self._render_block(self.doc, b, width)
        self._setup_footer()
        self._core_properties()
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


def render_docx(layout: LayoutDocument, path: Path) -> Path:
    return DocxRenderer(layout).render(path)
