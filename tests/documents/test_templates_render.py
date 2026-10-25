"""Every template × representative profiles → genuine DOCX and PDF that pass validation (TPL-01/02/06, DOC-01..03)."""

from __future__ import annotations

from datetime import date

import pytest

from ris.composition.composer import ComposeContext, ResumeComposer
from ris.domain.enums import RecordKind
from ris.rendering import extract
from ris.rendering.docx_renderer import render_docx
from ris.rendering.layout import plan_layout
from ris.rendering.libreoffice import find_soffice
from ris.rendering.pdf_renderer import PdfRenderer
from ris.rendering.specs import TemplateLibrary
from ris.rendering.validation import validate_outputs
from ris.services.profile import ProfileService
from ris.storage.workspace import Workspace
from tests import fixtures

LIB = TemplateLibrary()
TEMPLATE_IDS = [t.id for t in LIB.list()]
BUILDERS = {
    "short": fixtures.build_short,
    "incomplete": fixtures.build_incomplete,
    "long": fixtures.build_long,
    "unicode": fixtures.build_unicode,
}


@pytest.fixture(scope="module")
def profile_workspaces(tmp_path_factory):
    out = {}
    for name, builder in BUILDERS.items():
        ws = Workspace(tmp_path_factory.mktemp(f"p-{name}"), name, device_id="t").open()
        builder(ws.repo)
        out[name] = ws
    yield out
    for ws in out.values():
        ws.close()


def _mode_for(spec) -> str:
    return spec.recommended_modes[0] if spec.recommended_modes else "general_professional"


def _render(ws, spec, tmp_path, mode=None, **variant_extra):
    from ris.domain.models import Variant

    snap = ProfileService(ws.repo).snapshot()
    v = Variant(name="t", mode=mode or _mode_for(spec), template_id=spec.id, length="full", **variant_extra)
    ctx = ComposeContext(snapshot=snap, variant=v, template_id=spec.id, template_version=spec.version,
                         date_style=spec.entry.date_style, photo_allowed=spec.photo.allowed, today=date(2026, 10, 24))
    doc = ResumeComposer(ctx).compose()
    layout = plan_layout(doc, spec)
    docx = render_docx(layout, tmp_path / f"{spec.id}.docx")
    res = PdfRenderer(layout).render(tmp_path / f"{spec.id}.pdf")
    report = validate_outputs(doc, docx_path=docx, pdf_path=tmp_path / f"{spec.id}.pdf",
                              footer_strings=[layout.footer.left_text, layout.footer.note],
                              missing_glyphs=res.missing_glyphs)
    return doc, layout, docx, tmp_path / f"{spec.id}.pdf", res, report


@pytest.mark.req("TPL-01")
def test_twelve_distinct_layout_systems():
    specs = LIB.list()
    assert len(specs) >= 12 and not LIB.load_errors
    signatures = {(s.header.style, s.heading.style, s.entry.style, s.skills.style, s.columns.enabled) for s in specs}
    assert len(signatures) == len(specs), "templates must differ in layout primitives, not only colour"
    names = {s.name for s in specs}
    for expected in ["Executive Technical", "Defence Technical", "Scientific Research", "Engineering Professional",
                     "AI/ML Specialist", "Avionics & Drone Systems", "Government CV", "ATS Minimal",
                     "Modern Professional", "Academic Extended CV", "Compact Two-Page Resume", "Custom Diary Edition"]:
        assert expected in names


@pytest.mark.req("TPL-02")
def test_ats_template_is_plain(profile_workspaces, tmp_path):
    spec = LIB.get("ats_minimal")
    assert spec.ats_safe and not spec.tables_allowed and not spec.columns.enabled
    _, layout, docx, _, _, report = _render(profile_workspaces["long"], spec, tmp_path, mode="ats_plain")
    info = extract.read_docx(docx)
    assert info.tables == 0, "ATS template must not use tables"
    from ris.rendering.blocks import BannerBlock, ColumnsBlock, ImageBlock, LabeledSection, TableBlock

    assert not any(isinstance(b, (TableBlock, ColumnsBlock, BannerBlock, LabeledSection, ImageBlock)) for b in layout.blocks)
    assert report.ok, list(report.failures)


@pytest.mark.req("TPL-06", "DOC-01", "DOC-02", "DOC-03", "DOC-06")
@pytest.mark.parametrize("template_id", TEMPLATE_IDS)
@pytest.mark.parametrize("profile", list(BUILDERS))
def test_template_profile_matrix(profile_workspaces, tmp_path, template_id, profile):
    spec = LIB.get(template_id)
    doc, layout, docx, pdf, res, report = _render(profile_workspaces[profile], spec, tmp_path)
    failures = [(c.name, c.details[:3]) for c in report.failures]
    assert report.ok, failures
    info = extract.read_docx(docx)
    assert info.heading_paragraphs, "DOCX must use heading styles"
    pinfo = extract.read_pdf(pdf)
    assert pinfo.page_count == res.page_count >= 1
    assert not pinfo.unembedded_fonts
    w, h = pinfo.page_sizes[0]
    assert abs(w - 595.3) < 1 and abs(h - 841.9) < 1  # A4
    if profile == "long":
        assert pinfo.page_count >= 3
    if profile == "unicode":
        assert "Śarmā" in pinfo.text and "Σ-Δ" in pinfo.text


@pytest.mark.req("DOC-02")
def test_docx_structures(profile_workspaces, tmp_path):
    spec = LIB.get("academic_extended")
    _, _, docx, _, _, _ = _render(profile_workspaces["long"], spec, tmp_path, mode="academic_cv")
    info = extract.read_docx(docx)
    assert info.bullet_paragraphs > 20, "bullets must be real Word list paragraphs"
    assert info.has_page_field, "footer must contain a PAGE field"
    assert "https://orcid.org/0000-0002-1825-0097" in info.hyperlinks
    assert {"Title", "Heading 1"} <= info.styles_used
    from docx import Document

    d = Document(str(docx))
    assert d.core_properties.author == "Dr. Long Profile Example"
    assert "Resume Intelligence Studio" in d.core_properties.comments


@pytest.mark.req("DOC-03")
def test_pdf_metadata_links_outline(profile_workspaces, tmp_path):
    from pypdf import PdfReader

    spec = LIB.get("executive_technical")
    _, _, _, pdf, _, _ = _render(profile_workspaces["long"], spec, tmp_path)
    r = PdfReader(str(pdf))
    assert r.metadata.author == "Dr. Long Profile Example"
    assert "Resume Intelligence Studio" in (r.metadata.producer or "")
    assert r.outline, "section headings must produce PDF bookmarks"
    info = extract.read_pdf(pdf)
    assert "mailto:long.profile@example.org" in info.links


def test_government_mode_tables_and_declaration(profile_workspaces, tmp_path):
    spec = LIB.get("government_cv")
    doc, _, docx, pdf, _, report = _render(profile_workspaces["long"], spec, tmp_path, mode="government_biodata")
    assert report.ok, [c.details for c in report.failures]
    keys = [s.key for s in doc.sections]
    assert keys[0] == "particulars" and keys[-1] == "declaration"
    info = extract.read_docx(docx)
    assert info.tables >= 3
    assert "BIO-DATA / CURRICULUM VITAE" in info.body_text
    assert "I hereby declare" in extract.read_pdf(pdf).text


@pytest.mark.libreoffice
@pytest.mark.req("DOC-04")
@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
@pytest.mark.parametrize("template_id", ["defence_technical", "aiml_specialist", "modern_professional"])
def test_libreoffice_conversion_consistency(profile_workspaces, tmp_path, template_id):
    from ris.rendering.libreoffice import convert_docx_to_pdf

    spec = LIB.get(template_id)
    doc, layout, docx, _, _, _ = _render(profile_workspaces["long"], spec, tmp_path)
    lo_pdf = convert_docx_to_pdf(docx, tmp_path / "lo.pdf", find_soffice())
    report = validate_outputs(doc, docx_path=docx, pdf_path=lo_pdf,
                              footer_strings=[layout.footer.left_text, layout.footer.note])
    assert report.ok, [(c.name, c.details[:3]) for c in report.failures]
    from pypdf import PdfReader

    assert "/StructTreeRoot" in PdfReader(str(lo_pdf)).trailer["/Root"], "LibreOffice output should be tagged"


def test_variant_overrides_change_layout(profile_workspaces, tmp_path):
    spec = LIB.get("ats_minimal")
    base_doc, base_layout, *_ = _render(profile_workspaces["long"], spec, tmp_path, mode="ats_plain")
    _, small_layout, *_ = _render(profile_workspaces["long"], spec, tmp_path, mode="ats_plain", font_scale=0.9,
                                  spacing="compact", accent_color="#701C35", page_size="Letter")
    assert small_layout.styles["Body"].size < base_layout.styles["Body"].size
    assert abs(small_layout.page.width - 612) < 0.5
    del base_doc


def test_ws_kind_counts(profile_workspaces):
    assert profile_workspaces["long"].repo.count(RecordKind.PUBLICATION) == 40


@pytest.mark.req("DOC-06")
def test_unsupported_script_is_reported_not_hidden(tmp_path):
    """Characters the embedded fonts cannot draw must fail validation with a clear message."""
    ws = Workspace(tmp_path / "deva", "d", device_id="t").open()
    try:
        fixtures.build_unicode(ws.repo)
        rec = ws.repo.singleton(RecordKind.PERSONAL)
        ws.repo.patch(rec.id, preferred_name="Dr. Ānanya Śarmā (अनन्या)")
        spec = LIB.get("ats_minimal")
        _, _, _, _, res, report = _render(ws, spec, tmp_path)
        assert "अ" in res.missing_glyphs
        glyph = next(c for c in report.checks if c.name == "glyphs")
        assert glyph.status == "fail" and "LibreOffice" in glyph.message
        assert not report.ok
    finally:
        ws.close()


@pytest.mark.req("DOC-06")
def test_validation_readings_tolerate_pdf_extraction_artefacts():
    """Second-chance readings for text extracted from LibreOffice PDFs (the original text is always checked too)."""
    from ris.rendering.validation import _alt_reading, _letter_spaced_reading

    tok = extract.tokens
    assert tok(_letter_spaced_reading("A R E A S  O F  E X P E R T I S E\n▪ Avionics")) == \
        ["areas", "of", "expertise", "avionics"]
    assert tok(_letter_spaced_reading("E D U C A T I O NPh.D., Electrical"))[:2] == ["education", "ph.d"]
    assert _letter_spaced_reading("Normal text with words") == "Normal text with words"
    assert "funding" in tok(_alt_reading("F UNDING"))
